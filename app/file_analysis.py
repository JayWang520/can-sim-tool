"""Read-only, bounded offline CAN file import and statistics."""
import csv
import io
import json
import math
import re
import zipfile
from pathlib import Path

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100000
ALIASES = {
    'ts': ('ts', 'timestamp', 'time', '时间', '时间戳'),
    'id': ('id', 'can_id', 'arbitration_id', '帧id'),
    'data': ('data', 'data_hex', '数据'),
    'ext': ('ext', 'extended', 'is_extended_id', '扩展帧'),
    'dir': ('dir', 'direction', '方向'),
    'dlc': ('dlc', 'length', '长度'),
    'channel': ('channel', '通道'),
}


def analyze_file(name, content, sheet='', id_base=16, time_unit='s'):
    if len(content) > MAX_BYTES:
        raise ValueError('文件超过 20 MiB，请先拆分')
    if id_base not in (10, 16) or time_unit not in ('s', 'ms', 'us'):
        raise ValueError('无效的 ID 进制或时间单位')
    suffix = Path(name).suffix.lower()
    sheets = []
    workbook = None
    try:
        if suffix == '.xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError('XLSX 解压内容超过 100 MiB')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheets = workbook.sheetnames
            sheet = sheet or sheets[0]
            if sheet not in sheets:
                raise ValueError('工作表不存在')
            values = workbook[sheet].iter_rows(values_only=True)
            headers = next(values, ())
            rows = ((i, dict(zip(headers, row))) for i, row in enumerate(values, 2))
        elif suffix in ('.csv', '.tsv'):
            try:
                source = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                source = content.decode('gb18030')
            reader = csv.DictReader(io.StringIO(source), delimiter='\t' if suffix == '.tsv' else ',')
            rows = enumerate(reader, 2)
        elif suffix == '.json':
            records = json.loads(content)
            if isinstance(records, dict):
                records = records.get('frames')
            if not isinstance(records, list):
                raise ValueError('JSON 须为报文数组或 {"frames": [...]}')
            rows = enumerate(records, 1)
        elif suffix == '.asc':
            source = content.decode('utf-8-sig')
            if re.search(r'base\s+dec|timestamps\s+relative', source, re.I):
                raise ValueError('ASC 目前支持 base hex、timestamps absolute')
            def asc_rows():
                for i, line in enumerate(source.splitlines(), 1):
                    if not line.strip() or re.match(r'\s*(date|base|internal|//|begin|end)', line, re.I):
                        continue
                    match = re.match(r'^\s*(\S+)\s+(\S+)\s+([0-9a-f]+?)(x)?\s+(rx|tx)\s+d\s+(\d+)\s*(.*?)\s*$', line, re.I)
                    if not match:
                        yield i, None
                        continue
                    ts, channel, cid, ext, direction, dlc, data = match.groups()
                    yield i, dict(ts=ts, channel=channel, id=cid, ext=bool(ext), dir=direction, dlc=dlc, data=data)
            rows = asc_rows()
            id_base, time_unit = 16, 's'
        else:
            raise ValueError('支持 CSV、TSV、XLSX、JSON、ASC；不支持旧版 XLS')

        frames, errors = [], []
        total = 0
        for line, row in rows:
            total += 1
            if total > MAX_ROWS:
                raise ValueError('超过 100000 行，请拆分文件；未返回不完整分析')
            try:
                if not isinstance(row, dict):
                    raise ValueError('不是有效报文记录（ASC 仅支持经典 CAN 数据帧）')
                keys = {str(k).strip().lower(): v for k, v in row.items()}
                r = {key: next((keys[a] for a in aliases if a in keys), None) for key, aliases in ALIASES.items()}
                if r['ts'] is None or r['id'] is None or r['data'] is None:
                    raise ValueError('缺少 ts、id 或 data 列/值（零字节 data 请用空字符串）')
                ts = float(r['ts']) / {'s': 1, 'ms': 1000, 'us': 1000000}[time_unit]
                if not math.isfinite(ts):
                    raise ValueError('时间必须是有限数值')
                raw_id = str(r['id']).strip()
                cid = int(raw_id, 16 if raw_id.lower().startswith('0x') else id_base)
                ext_value = str(r['ext']).strip().lower()
                if r['ext'] is None or ext_value == '':
                    ext = cid > 0x7ff
                elif ext_value in ('true', '1', 'extended', '扩展帧'):
                    ext = True
                elif ext_value in ('false', '0', 'standard', '标准帧'):
                    ext = False
                else:
                    raise ValueError('ext 须为 0/1 或 false/true')
                if not 0 <= cid <= (0x1fffffff if ext else 0x7ff):
                    raise ValueError('ID 超出帧类型范围')
                if isinstance(r['data'], list):
                    if any(type(v) is not int or not 0 <= v <= 255 for v in r['data']):
                        raise ValueError('字节数组须为 0–255 整数')
                    data = bytes(r['data'])
                elif isinstance(r['data'], str):
                    data = bytes.fromhex(r['data'])
                else:
                    raise ValueError('data 须为十六进制文本；Excel 请将数据列设为文本')
                if len(data) > 8:
                    raise ValueError('目前只支持经典 CAN，最多 8 字节')
                if r['dlc'] not in (None, '') and int(str(r['dlc'])) != len(data):
                    raise ValueError('DLC 与数据字节数不一致')
                direction = str(r['dir'] or 'rx').lower()
                if direction not in ('rx', 'tx'):
                    raise ValueError('方向须为 rx/tx')
                frames.append(dict(row=line, ts=ts, id=f'{cid:08X}' if ext else f'{cid:03X}', ext=ext,
                                   dir=direction, channel=str(r['channel'] or ''), dlc=len(data), data=data.hex().upper()))
            except (ValueError, TypeError, OverflowError) as exc:
                errors.append(dict(row=line, error=str(exc)))
        groups = {}
        for f in frames:
            key = (f['channel'], f['id'], f['ext'], f['dir'])
            groups.setdefault(key, []).append(f)
        stats = []
        for (channel, cid, ext, direction), group in groups.items():
            intervals = [b['ts'] - a['ts'] for a, b in zip(group, group[1:])]
            valid = [v * 1000 for v in intervals if v >= 0]
            byte_stats = []
            for index in range(8):
                values = [int(f['data'][index*2:index*2+2], 16) if f['dlc'] > index else None for f in group]
                present = [v for v in values if v is not None]
                byte_stats.append(dict(byte=index+1, min=min(present) if present else None,
                                       max=max(present) if present else None,
                                       changes=sum(a != b for a, b in zip(values, values[1:]) if a is not None and b is not None)))
            stats.append(dict(channel=channel, id=cid, ext=ext, dir=direction, count=len(group),
                              period_min_ms=min(valid) if valid else None, period_max_ms=max(valid) if valid else None,
                              period_mean_ms=sum(valid)/len(valid) if valid else None,
                              backwards=sum(v < 0 for v in intervals),
                              changes=sum(a['data'] != b['data'] for a, b in zip(group, group[1:])), bytes=byte_stats))
        return dict(file=name, sheet=sheet, sheets=sheets, total=total, count=len(frames), errors=errors,
                    duration_s=max(f['ts'] for f in frames)-min(f['ts'] for f in frames) if frames else 0,
                    stats=stats, frames=frames)
    finally:
        if workbook is not None:
            workbook.close()
