import io
import json

import pytest
from openpyxl import Workbook
from fastapi.testclient import TestClient
from app.file_analysis import analyze_file
from app.main import app


ROWS = [dict(ts=0, id='123', extended=0, dir='rx', dlc=2, data='1234'),
        dict(ts=0.01, id='123', extended=0, dir='rx', dlc=2, data='1256')]


@pytest.mark.parametrize('extension', ['csv', 'tsv', 'json', 'xlsx', 'asc'])
def test_formats(extension):
    if extension == 'json':
        content = json.dumps(ROWS).encode()
    elif extension == 'xlsx':
        wb = Workbook()
        wb.active.append(list(ROWS[0]))
        for row in ROWS:
            wb.active.append(list(row.values()))
        wb.create_sheet('Other')
        stream = io.BytesIO(); wb.save(stream); content = stream.getvalue()
    elif extension == 'asc':
        content = b'base hex timestamps absolute\n0 1 123 Rx d 2 12 34\n0.01 1 123 Rx d 2 1256\n'
    else:
        sep = '\t' if extension == 'tsv' else ','
        content = ('\n'.join([sep.join(ROWS[0])] + [sep.join(map(str, row.values())) for row in ROWS])).encode()
    with TestClient(app) as client:
        response = client.post('/api/files/analyze', files={'file': (f'frames.{extension}', content)})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['count'] == 2 and not result['errors']
    group = result['stats'][0]
    assert group['period_mean_ms'] == pytest.approx(10)
    assert group['bytes'][1] == dict(byte=2, min=52, max=86, changes=1)
    assert result['frames'][0]['data'] == '1234'


def test_invalid_rows_and_identity():
    rows = ROWS + [dict(ROWS[0], extended=1), dict(ROWS[0], channel='2'),
                   dict(ROWS[0], data='1'), dict(ROWS[0], dlc=8),
                   dict(ROWS[0], ts='nan'), dict(ROWS[0], extended='nope'),
                   dict(ROWS[0], id='800'), dict(ROWS[0], data=1234)]
    result = analyze_file('a.json', json.dumps(rows).encode())
    assert result['count'] == 4
    assert len(result['stats']) == 3
    assert [e['row'] for e in result['errors']] == list(range(5, 11))


def test_units_backwards_and_empty():
    result = analyze_file('a.json', json.dumps([dict(ROWS[0], ts=20), dict(ROWS[1], ts=10)]).encode(), id_base=10, time_unit='ms')
    assert result['frames'][0]['id'] == '07B'
    assert result['stats'][0]['backwards'] == 1
    assert result['stats'][0]['period_mean_ms'] is None
    assert analyze_file('a.json', b'[]')['count'] == 0


def test_limits_and_unsupported(monkeypatch):
    import app.file_analysis as module
    monkeypatch.setattr(module, 'MAX_ROWS', 1)
    with pytest.raises(ValueError, match='100000'):
        analyze_file('a.json', json.dumps(ROWS).encode())
    with pytest.raises(ValueError):
        analyze_file('a.xls', b'xxx')
    with pytest.raises(ValueError):
        analyze_file('a.asc', b'base dec timestamps relative')


def test_import_never_sends(monkeypatch):
    from app.main import bus
    monkeypatch.setattr(bus, 'send', lambda *a, **kw: pytest.fail('offline import sent a frame'))
    with TestClient(app) as client:
        assert client.post('/api/files/analyze', files={'file': ('a.json', json.dumps(ROWS))}).status_code == 200
        assert client.post('/api/files/analyze', files={'file': ('a.xlsx', b'broken')}).status_code == 400


def test_xlsx_sheet_and_formula_preservation():
    wb = Workbook()
    wb.active.title = 'Empty'
    ws = wb.create_sheet('CAN')
    ws.append(list(ROWS[0]))
    ws.append(list(ROWS[0].values()))
    ws.append(['=1+1', '123', 0, 'rx', 2, '1234'])
    stream = io.BytesIO(); wb.save(stream)
    content = stream.getvalue()
    result = analyze_file('a.xlsx', content, sheet='CAN')
    assert result['sheets'] == ['Empty', 'CAN']
    assert result['count'] == 1
    assert result['errors'][0]['row'] == 3
    assert stream.getvalue() == content
    with pytest.raises(ValueError, match='工作表不存在'):
        analyze_file('a.xlsx', content, sheet='missing')
